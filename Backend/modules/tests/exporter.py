import io
import os
import docx
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fpdf import FPDF

# --- КОНСТАНТЫ ---
UNIVERSITY_INFO = (
    "ФГБОУ ВО «Кубанский государственный университет»\n"
    "Кафедра вычислительных технологий"
)
CHIEF_TITLE = "И. о. заведующего кафедрой вычислительных технологий"
CHIEF_NAME = "Т.А. Приходько"
FONT_NAME = "Times New Roman"

# --- КЛАСС PDF ---
class PDF(FPDF):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        base_dir = os.getcwd()
        font_path = lambda f: os.path.join(base_dir, "static", "fonts", f)
        
        try:
            self.add_font('TimesNewRoman', '', font_path("TIMES.TTF"), uni=True)
            self.add_font('TimesNewRoman', 'B', font_path("TIMESBD.TTF"), uni=True)
        except Exception as e:
            raise FileNotFoundError(f"Критическая ошибка: шрифты не найдены. {e}")

def get_ticket_header(discipline_name):
    return f"{UNIVERSITY_INFO}\nДисциплина «{discipline_name}»"

# Генерация DOCX
def generate_tests_docx(discipline_name: str, tests: list):
    doc = Document()
    
    section = doc.sections[0]
    section.left_margin, section.right_margin = Cm(2), Cm(1)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)

    def add_styled_para(text, size, bold=False, align=None, keep=True, space_after=0, space_before=0, indent=0):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.bold = bold
        if align: p.alignment = align
        if keep:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if indent: p.paragraph_format.first_line_indent = Cm(indent)
        return p

    for i, test in enumerate(tests):
        # Шапка
        add_styled_para(get_ticket_header(discipline_name), 11, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Номер билета
        add_styled_para(f"БИЛЕТ № {test.test_number}", 12, bold=True, 
                        align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=6)

        # Вопросы
        for idx, question in enumerate(test.questions, 1):
            add_styled_para(f"{idx}. {question.question_content}", 12, 
                            align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, space_after=2)

        # Прослойка
        add_styled_para("\n", 10, space_after=0, space_before=0)

        # Подпись (Таблица)
        table = doc.add_table(rows=1, cols=2)
        table.width = Cm(18)
        table.left_padding = table.right_padding = 0
        table.columns[0].width = Cm(14.5)
        table.columns[1].width = Cm(3.5)

        trPr = table.rows[0]._tr.get_or_add_trPr()
        cantSplit = docx.oxml.shared.OxmlElement('w:cantSplit')
        cantSplit.set(qn('w:val'), 'true')
        trPr.append(cantSplit)

        # Ячейка должности
        c1_p = table.cell(0, 0).paragraphs[0]
        c1_p._element.get_or_add_pPr().append(docx.oxml.shared.OxmlElement('w:suppressHyphenation'))
        run_l = c1_p.add_run(CHIEF_TITLE)
        run_l.font.name, run_l.font.size = FONT_NAME, Pt(11)
        
        # Ячейка фамилии
        c2_p = table.cell(0, 1).paragraphs[0]
        c2_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_r = c2_p.add_run(CHIEF_NAME)
        run_r.font.name, run_r.font.size = FONT_NAME, Pt(11)

        if i < len(tests) - 1:
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream

# Генерация PDF
def generate_tests_pdf(discipline_name: str, tests: list):
    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    def setup_font(style='', size=11):
        pdf.set_font('TimesNewRoman', style, size)

    pdf.add_page()
    
    for i, test in enumerate(tests):
        if pdf.get_y() > 210:
            pdf.add_page()

        # Шапка
        setup_font(size=11)
        pdf.multi_cell(0, 5, get_ticket_header(discipline_name), align='C')
        
        # Заголовок
        pdf.ln(2)
        setup_font('B', 12)
        pdf.cell(0, 8, f"БИЛЕТ № {test.test_number}", align='C', ln=True)

        # Вопросы
        setup_font(size=12)
        for idx, q in enumerate(test.questions, 1):
            pdf.set_x(20)
            pdf.multi_cell(0, 6, f"{idx}. {q.question_content}", align='J')
        
        # Подпись
        pdf.ln(5)
        setup_font(size=11)
        pdf.cell(145, 5, CHIEF_TITLE, align='L')
        pdf.cell(0, 5, CHIEF_NAME, align='R', ln=True)

        if i < len(tests) - 1:
            pdf.ln(10)

    return io.BytesIO(pdf.output())