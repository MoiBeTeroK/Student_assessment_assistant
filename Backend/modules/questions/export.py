import io
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

# Константы
HEADER_TEXT = (
    "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ\n"
    "ФГБОУ ВО «Кубанский государственный университет»\n"
    "Кафедра вычислительных технологий\n"
)
FONT_NAME_DOCX = 'Times New Roman'
FONT_NAME_PDF = 'TimesNewRoman'


class PDF(FPDF):
    """
    Кастомный класс FPDF с автоматической загрузкой кириллических 
    TrueType шрифтов из директории проекта при инициализации.
    """
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        base_dir = os.getcwd()
        font_path = lambda f: os.path.join(base_dir, "static", "fonts", f)
        
        try:
            self.add_font(FONT_NAME_PDF, '', font_path("TIMES.ttf"))
            self.add_font(FONT_NAME_PDF, 'B', font_path("TIMESBD.ttf"))
        except Exception as e:
            print(f"Предупреждение: кириллические шрифты не загружены ({e})")


def get_title(discipline_name: str) -> str:
    """Формирует стандартизированный текст заголовка для документа."""
    return f"Вопросы к экзамену по дисциплине\n«{discipline_name}»"


def generate_docx(discipline_name: str, questions: list) -> io.BytesIO:
    doc = Document()
    
    # Настройка полей страницы
    section = doc.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(1)

    def add_para(text, size=12, bold=False, align=None, indent=0, space_after=0):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = FONT_NAME_DOCX
        run.font.size = Pt(size)
        run.bold = bold
        if align: 
            p.alignment = align
        if indent: 
            p.paragraph_format.first_line_indent = Cm(indent)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.0
        return p

    # Добавление шапки организации
    add_para(HEADER_TEXT, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Добавление заголовка документа
    add_para(get_title(discipline_name), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # Добавление пронумерованного списка вопросов с абзацным отступом
    for i, q_text in enumerate(questions, 1):
        add_para(f"{i}. {q_text}", indent=1.25, space_after=2)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


def generate_pdf(discipline_name: str, questions: list) -> io.BytesIO:
    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Отрисовка шапки организации
    pdf.set_font(FONT_NAME_PDF, 'B', 11)
    pdf.multi_cell(0, 4, HEADER_TEXT, align="C")
    pdf.ln(3) 
    
    # Отрисовка заголовка документа
    pdf.multi_cell(0, 5, get_title(discipline_name), align="C")
    pdf.ln(4)

    # Отрисовка вопросов
    pdf.set_font(FONT_NAME_PDF, '', 12)
    
    # Настройка абзацного отступа (12.5 мм = 1.25 см)
    indent_width = 12.5
    
    # Расчет чистой доступной ширины текстового блока
    effective_page_width = pdf.w - pdf.l_margin - pdf.r_margin
    usable_text_width = effective_page_width - indent_width
    
    for i, q_text in enumerate(questions, 1):
        # Сдвигаем курсор X на величину левого поля + абзацный отступ
        pdf.set_x(pdf.l_margin + indent_width) 
        
        # Передаем строго вычисленную ширину ячейки во избежание обрезки справа
        pdf.multi_cell(usable_text_width, 5, f"{i}. {q_text}")
        pdf.ln(2)

    # Генерация бинарного контента документа
    pdf_bytes = pdf.output()
    
    # Безопасное приведение структуры fpdf2 к сырому потоку байт
    if isinstance(pdf_bytes, str):
        pdf_bytes = pdf_bytes.encode('utf-8')
    elif isinstance(pdf_bytes, bytearray):
        pdf_bytes = bytes(pdf_bytes)
        
    return io.BytesIO(pdf_bytes)
